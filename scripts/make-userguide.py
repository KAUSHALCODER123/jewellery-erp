"""Render USER_GUIDE.md to USER_GUIDE.pdf and USER_GUIDE.docx.

A small Markdown-subset parser (headings, paragraphs, bullet/numbered lists,
tables, blockquotes, horizontal rules, inline **bold** and `code`) feeds two
renderers: fpdf2 for PDF and python-docx for Word. Currency rupee signs are
converted to "Rs" so the PDF renders with core fonts (no asset bundling).
"""

import os
import re

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "USER_GUIDE.md")

NAVY = (0x1F, 0x2E, 0x4E)
ACCENT = (0xB8, 0x86, 0x0B)


def clean(text):
    # Replace non-latin-1 glyphs with PDF-core-font-safe equivalents, then drop
    # anything still outside latin-1 (e.g. emoji) so core Helvetica can render it.
    repl = {
        "₹": "Rs ", "→": "->", "–": "-", "—": "-",
        "‘": "'", "’": "'", "“": '"', "”": '"', "×": "x",
        "≥": ">=", "≤": "<=", "✓": "[ok]", "✅": "", "⚠️": "!", "⚠": "!",
        "•": chr(149), "·": "-", "…": "...", "®": "(R)", "™": "(TM)",
        "“": '"', "🔴": "", "🆕": "(new)",
    }
    for k, v in repl.items():
        text = text.replace(k, v)
    return text.encode("latin-1", "ignore").decode("latin-1")


def strip_md_links(text):
    # [label](url) -> label
    return re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)


def parse(md):
    """Return a list of block dicts."""
    lines = md.split("\n")
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            blocks.append({"type": "hr"})
            i += 1
            continue

        m = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if m:
            blocks.append({"type": "h", "level": len(m.group(1)), "text": m.group(2)})
            i += 1
            continue

        # Table: a line with | followed by a |---| separator
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            rows = []
            header = [c.strip() for c in stripped.strip("|").split("|")]
            i += 2  # skip header + separator
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            blocks.append({"type": "table", "header": header, "rows": rows})
            continue

        if stripped.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append({"type": "quote", "text": " ".join(quote)})
            continue

        # List item (bullet or numbered), with indent -> nesting level
        lm = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if lm:
            items = []
            while i < len(lines):
                im = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", lines[i])
                if not im:
                    if lines[i].strip() == "":
                        break
                    # continuation line of previous item
                    if items:
                        items[-1]["text"] += " " + lines[i].strip()
                        i += 1
                        continue
                    break
                indent = len(im.group(1))
                ordered = bool(re.match(r"\d+\.", im.group(2)))
                items.append({"level": 0 if indent < 2 else 1, "ordered": ordered, "text": im.group(3)})
                i += 1
            blocks.append({"type": "list", "items": items})
            continue

        # Paragraph: gather consecutive plain lines
        para = [stripped]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,3}\s|>|\||\s*([-*]|\d+\.)\s|---$)", lines[i].strip()):
            para.append(lines[i].strip())
            i += 1
        blocks.append({"type": "p", "text": " ".join(para)})
    return blocks


# ----------------------------- PDF -----------------------------

def build_pdf(blocks, path):
    from fpdf import FPDF
    from fpdf.fonts import FontFace

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.set_margins(16, 15, 16)
    pdf.add_page()
    width = 210 - 32

    def md_text(t):
        return clean(strip_md_links(t))

    for b in blocks:
        t = b.get("type")
        if t == "hr":
            pdf.set_draw_color(210, 210, 210)
            pdf.set_line_width(0.2)
            y = pdf.get_y() + 1
            pdf.line(16, y, 210 - 16, y)
            pdf.ln(3)
        elif t == "h":
            lvl = b["level"]
            pdf.ln(2 if lvl > 1 else 3)
            pdf.set_text_color(*NAVY)
            size = {1: 20, 2: 14, 3: 11}[lvl]
            pdf.set_font("Helvetica", "B", size)
            pdf.multi_cell(width, size * 0.5, md_text(b["text"]), markdown=True)
            if lvl == 1:
                pdf.set_draw_color(*ACCENT)
                pdf.set_line_width(0.6)
                pdf.line(16, pdf.get_y() + 1, 210 - 16, pdf.get_y() + 1)
                pdf.ln(2)
            pdf.ln(1)
        elif t == "p":
            pdf.set_text_color(20, 20, 20)
            pdf.set_font("Helvetica", "", 9.5)
            pdf.multi_cell(width, 4.8, md_text(b["text"]), markdown=True)
            pdf.ln(1.5)
        elif t == "quote":
            pdf.set_text_color(90, 70, 10)
            pdf.set_font("Helvetica", "I", 9)
            pdf.set_x(20)
            pdf.multi_cell(width - 4, 4.6, md_text(b["text"]), markdown=True)
            pdf.ln(1.5)
        elif t == "list":
            pdf.set_text_color(20, 20, 20)
            pdf.set_font("Helvetica", "", 9.5)
            for it in b["items"]:
                indent = 4 + it["level"] * 6
                bullet = chr(149) + " " if not it["ordered"] else "- "
                pdf.set_x(16 + indent)
                pdf.multi_cell(width - indent, 4.6, bullet + md_text(it["text"]), markdown=True)
            pdf.ln(1.5)
        elif t == "table":
            pdf.set_font("Helvetica", "", 8.5)
            pdf.set_text_color(20, 20, 20)
            with pdf.table(width=width, text_align="LEFT", line_height=4.6,
                           headings_style=FontFace(emphasis="BOLD", fill_color=(230, 235, 242))) as table:
                row = table.row()
                for c in b["header"]:
                    row.cell(md_text(c))
                for r in b["rows"]:
                    row = table.row()
                    for c in r:
                        row.cell(md_text(c))
            pdf.ln(2)

    pdf.output(path)
    print("PDF written:", path)


# ----------------------------- DOCX -----------------------------

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_runs(paragraph, text, base_size):
    from docx.shared import Pt
    text = strip_md_links(text)
    for seg in INLINE_RE.split(text):
        if not seg:
            continue
        if seg.startswith("**") and seg.endswith("**"):
            r = paragraph.add_run(seg[2:-2]); r.bold = True
        elif seg.startswith("`") and seg.endswith("`"):
            r = paragraph.add_run(seg[1:-1]); r.font.name = "Consolas"
        else:
            r = paragraph.add_run(seg)
        r.font.size = Pt(base_size)


def build_docx(blocks, path):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.7); s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.8); s.right_margin = Inches(0.8)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    for b in blocks:
        t = b.get("type")
        if t == "hr":
            p = doc.add_paragraph()
            pr = p._p.get_or_add_pPr()
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "CCCCCC")
            pbdr.append(bottom); pr.append(pbdr)
        elif t == "h":
            lvl = b["level"]
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10 if lvl == 1 else 8)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(strip_md_links(b["text"]))
            r.bold = True
            r.font.size = Pt({1: 20, 2: 15, 3: 12}[lvl])
            r.font.color.rgb = RGBColor(*NAVY)
        elif t == "p":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, b["text"], 10.5)
        elif t == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, b["text"], 10)
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(90, 70, 10)
        elif t == "list":
            for it in b["items"]:
                style = "List Number" if it["ordered"] else "List Bullet"
                p = doc.add_paragraph(style=style)
                p.paragraph_format.left_indent = Inches(0.3 + it["level"] * 0.3)
                p.paragraph_format.space_after = Pt(2)
                add_runs(p, it["text"], 10.5)
        elif t == "table":
            ncols = len(b["header"])
            table = doc.add_table(rows=1, cols=ncols)
            table.style = "Light Grid Accent 1"
            for j, c in enumerate(b["header"]):
                cell = table.rows[0].cells[j]
                cell.paragraphs[0].add_run(strip_md_links(c)).bold = True
            for r in b["rows"]:
                cells = table.add_row().cells
                for j in range(ncols):
                    add_runs(cells[j].paragraphs[0], r[j] if j < len(r) else "", 9.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    doc.save(path)
    print("DOCX written:", path)


if __name__ == "__main__":
    with open(SRC, encoding="utf-8") as f:
        blocks = parse(f.read())
    build_pdf(blocks, os.path.join(ROOT, "USER_GUIDE.pdf"))
    build_docx(blocks, os.path.join(ROOT, "USER_GUIDE.docx"))
