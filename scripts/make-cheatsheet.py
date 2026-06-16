"""Generate a one-page POS counter cheat-sheet in both .docx and .pdf.

Single content source -> two renderers (python-docx + fpdf2). Currency is
written as "Rs" so it renders in PDF core fonts; no external assets needed.
"""

TITLE = "POS Counter Quick Reference"
SUBTITLE = "Jewellery ERP  -  keep this at the till"

# Each section: (heading, [lines]). A line starting with a digit+'. ' is a step.
SECTIONS = [
    ("1. Start of Day", [
        "Open the app and log in.",
        "Dashboard > Daily Rates: enter today's Gold 24K / 22K / 18K / Silver, then Save & Update Rates.",
        "Rates drive EVERY bill price - never skip this.",
    ]),
    ("2. Make a Sale  (POS Billing)", [
        "Columns: Customer (left)  |  Cart (center)  |  Payment (right).",
        "1. Customer: search by phone/name; or type a name + Enter for a walk-in; or click + for a new customer.",
        "2. Add items: scan the barcode (auto-adds).",
        "   - Scanner dead / tag damaged? Type the barcode or HUID + Enter.",
        "   - Item never tagged? Click Quick Bill > Weight-wise (gold needs HUID) or Flat price.",
        "3. Old-gold trade-in: URD panel > Description / Tunch / Wt / Rate > Add (value is deducted).",
        "4. Payment: split across Cash / UPI / Card / Cheque / NEFT / Udhari. Balance must reach Rs 0.",
        "5. Checkout (F8) > print A4 (GST) / A5 / Thermal, or Send via WhatsApp.",
        "Estimate only (no stock reduced)? Build the cart, then Save as Quotation.",
    ]),
    ("3. Must-Know Rules", [
        "GOLD sells only if hallmarked - enter the 6-char HUID stamped on the piece.",
        "Cash >= Rs 2,00,000 OR any old-gold exchange -> red panel needs PAN + Aadhaar + ID upload.",
        "Lock the screen when you leave the counter (auto-locks after 15 min).",
        "App works offline; only live rates and WhatsApp need internet.",
    ]),
    ("4. Quick Error Fixes", [
        "'No item found for ...' -> wrong/typo barcode; retype, or use Quick Bill.",
        "'... is sold ...' -> already billed at another counter; rescan inventory.",
        "'... has not been hallmarked (HUID is required)' -> gold needs its HUID; enter it, or bill as Flat price.",
        "Checkout greyed out -> Balance Remaining is not Rs 0; adjust a payment.",
        "'Overpaid - reduce a payment' -> lower one of the payment amounts.",
    ]),
    ("5. End of Day", [
        "Day Book: match the cash drawer to today's transactions.",
        "Keep 'Backup on Exit' enabled, then close the app. Logout to end your shift.",
    ]),
]

FOOTER = "Full guide: USER_GUIDE.md   |   Screens you live in: Dashboard . POS Billing . Inventory . CRM . Day Book"

NAVY = (0x1F, 0x2E, 0x4E)
ACCENT = (0xB8, 0x86, 0x0B)  # gold/amber


def build_pdf(path):
    from fpdf import FPDF

    pdf = FPDF(orientation="P", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=False)
    pdf.set_margins(12, 10, 12)
    pdf.add_page()
    width = 210 - 24

    pdf.set_text_color(*NAVY)
    pdf.set_font("Helvetica", "B", 17)
    pdf.cell(0, 8, TITLE, new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "I", 9)
    pdf.set_text_color(90, 90, 90)
    pdf.cell(0, 5, SUBTITLE, new_x="LMARGIN", new_y="NEXT")
    pdf.set_draw_color(*ACCENT)
    pdf.set_line_width(0.6)
    y = pdf.get_y() + 1
    pdf.line(12, y, 210 - 12, y)
    pdf.ln(3)

    for heading, lines in SECTIONS:
        pdf.set_text_color(*NAVY)
        pdf.set_font("Helvetica", "B", 10.5)
        pdf.multi_cell(width, 5.0, heading, new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(20, 20, 20)
        pdf.set_font("Helvetica", "", 8.4)
        for line in lines:
            indent = 0
            text = line
            if line.startswith("   "):
                indent = 5
                text = line.strip()
            elif line[:2].rstrip().rstrip(".").isdigit() is False:
                # bullet for non-numbered, non-indented lines
                if not (len(line) > 1 and line[0].isdigit() and line[1] == "."):
                    text = chr(149) + " " + line  # bullet (cp1252 0x95)
            if indent:
                pdf.set_x(12 + indent)
            pdf.multi_cell(width - indent, 4.4, text, new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1.6)

    pdf.set_y(-16)
    pdf.set_draw_color(200, 200, 200)
    pdf.set_line_width(0.3)
    pdf.line(12, pdf.get_y(), 210 - 12, pdf.get_y())
    pdf.ln(1)
    pdf.set_font("Helvetica", "I", 7.5)
    pdf.set_text_color(110, 110, 110)
    pdf.multi_cell(width, 4, FOOTER, new_x="LMARGIN", new_y="NEXT")

    pdf.output(path)
    print("PDF written:", path)


def build_docx(path):
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(9)

    title = doc.add_paragraph()
    run = title.add_run(TITLE)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(*NAVY)
    sub = doc.add_paragraph()
    r2 = sub.add_run(SUBTITLE)
    r2.italic = True
    r2.font.size = Pt(9)
    r2.font.color.rgb = RGBColor(90, 90, 90)
    sub.paragraph_format.space_after = Pt(6)

    for heading, lines in SECTIONS:
        h = doc.add_paragraph()
        hr = h.add_run(heading)
        hr.bold = True
        hr.font.size = Pt(11)
        hr.font.color.rgb = RGBColor(*NAVY)
        h.paragraph_format.space_before = Pt(4)
        h.paragraph_format.space_after = Pt(2)
        for line in lines:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            text = line
            if line.startswith("   "):
                p.paragraph_format.left_indent = Inches(0.35)
                text = line.strip()
            elif not (len(line) > 1 and line[0].isdigit() and line[1] == "."):
                text = "• " + line
                p.paragraph_format.left_indent = Inches(0.18)
            else:
                p.paragraph_format.left_indent = Inches(0.18)
            run = p.add_run(text)
            run.font.size = Pt(8.8)

    f = doc.add_paragraph()
    f.paragraph_format.space_before = Pt(8)
    fr = f.add_run(FOOTER)
    fr.italic = True
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(110, 110, 110)

    doc.save(path)
    print("DOCX written:", path)


if __name__ == "__main__":
    import os
    out = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    build_pdf(os.path.join(out, "POS-Counter-Cheat-Sheet.pdf"))
    build_docx(os.path.join(out, "POS-Counter-Cheat-Sheet.docx"))
